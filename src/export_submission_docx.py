"""Render the authoritative Markdown manuscript to a review-ready DOCX with editable tables and embedded figures."""
from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"manuscript/PlantShiftEval_corrected_manuscript.md"
OUTPUT=ROOT/"manuscript/PlantShiftEval_corrected_manuscript.docx"
TITLE="Testing texture-representation hypotheses in plant disease recognition: training-strategy effects and cross-domain generalization limits"
FIGURES={
  1:ROOT/"report/corrected_summary/Figure1_source_comparisons.png",
  2:ROOT/"report/corrected_summary/Figure2_training_factorial.png",
  3:ROOT/"report/corrected_summary/Figure3_component_ablations.png",
  4:ROOT/"report/corrected_summary/Figure4_domain_limits.png",
  5:ROOT/"report/corrected_summary/Figure5_target_access.png",
}

def clean_inline(text):
    text=re.sub(r"\[([^\]]+)\]\(([^)]+)\)",r"\1 (\2)",text)
    text=re.sub(r"\*\*(.*?)\*\*",r"\1",text)
    text=re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)",r"\1",text)
    text=text.replace(chr(96),"")
    return text.strip()

def cells(line):
    return [x.strip() for x in line.strip().strip("|").split("|")]

def is_separator(line):
    vals=cells(line)
    return bool(vals) and all(re.fullmatch(r":?-{3,}:?",v or "-") for v in vals)

def set_font(run,name="Times New Roman",size=12,bold=None,italic=None):
    run.font.name=name;run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    rpr=run._element.get_or_add_rPr()
    fonts=rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts=OxmlElement("w:rFonts");rpr.append(fonts)
    for key in ("w:ascii","w:hAnsi","w:eastAsia","w:cs"):fonts.set(qn(key),name)

def add_page_field(paragraph):
    run=paragraph.add_run()
    begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText");instr.set(qn("xml:space"),"preserve");instr.text=" PAGE "
    sep=OxmlElement("w:fldChar");sep.set(qn("w:fldCharType"),"separate")
    text=OxmlElement("w:t");text.text="1"
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for item in (begin,instr,sep,text,end):run._r.append(item)

def add_table(doc,lines):
    parsed=[cells(line) for line in lines if line.strip() and not is_separator(line)]
    if not parsed:return
    width=max(len(row) for row in parsed)
    table=doc.add_table(rows=0,cols=width)
    table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ridx,row in enumerate(parsed):
        row=row+[""]*(width-len(row))
        cells_out=table.add_row().cells
        for cidx,value in enumerate(row):
            cell=cells_out[cidx];cell.text=clean_inline(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.line_spacing=1.0
                for run in par.runs:set_font(run,size=8.5,bold=(ridx==0))
            if ridx==0:
                tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"EAF2F8");tcPr.append(shd)
    doc.add_paragraph()

def main():
    if not SOURCE.exists():raise FileNotFoundError(SOURCE)
    lines=SOURCE.read_text(encoding="utf-8").splitlines()
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(1);sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after=Pt(0)
    for style_name in ("Heading 1","Heading 2","Heading 3"):
        st=doc.styles[style_name];st.font.name="Times New Roman";st.font.size=Pt(12);st.font.bold=True
        st.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(4)
    title=doc.add_paragraph()
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after=Pt(18)
    set_font(title.add_run(TITLE),size=16,bold=True)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_page_field(footer)
    sectPr=sec._sectPr;ln=OxmlElement("w:lnNumType");ln.set(qn("w:countBy"),"1");ln.set(qn("w:restart"),"continuous");sectPr.append(ln)
    i=0;in_refs=False;in_figures=False
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.lstrip().startswith("|"):
            block=[]
            while i<len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i]);i+=1
            add_table(doc,block);continue
        h=re.match(r"^(#{1,4})\s+(.+)$",line)
        if h:
            level=len(h.group(1));heading=clean_inline(h.group(2))
            in_refs=heading.lower()=="references";in_figures=heading.lower()=="figure legends"
            p=doc.add_heading(heading,level=min(level,3))
            for run in p.runs:set_font(run,size=12,bold=True)
            i+=1;continue
        if in_figures:
            fm=re.match(r"^\*\*Figure\s+(\d+)\.",line)
            if fm:
                caption=doc.add_paragraph()
                caption.paragraph_format.keep_with_next=True
                caption.paragraph_format.space_before=Pt(8)
                set_font(caption.add_run(clean_inline(line)),size=10,bold=False)
                fig=FIGURES.get(int(fm.group(1)))
                if fig and fig.exists():
                    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_together=True
                    p.add_run().add_picture(str(fig),width=Inches(6.1))
                else:
                    note=doc.add_paragraph(f"[Figure {fm.group(1)} image not found in repository]")
                    note.alignment=WD_ALIGN_PARAGRAPH.CENTER
                i+=1;continue
        if line.startswith(chr(96)*3):
            i+=1;code=[]
            while i<len(lines) and not lines[i].startswith(chr(96)*3):
                code.append(lines[i]);i+=1
            if i<len(lines):i+=1
            p=doc.add_paragraph("\n".join(code))
            for run in p.runs:set_font(run,name="Consolas",size=9)
            continue
        if re.match(r"^\s*[-*]\s+",line) and not in_refs:
            p=doc.add_paragraph(clean_inline(re.sub(r"^\s*[-*]\s+","",line)),style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+",line) and not in_refs:
            p=doc.add_paragraph(clean_inline(re.sub(r"^\s*\d+\.\s+","",line)),style="List Number")
        else:
            p=doc.add_paragraph(clean_inline(line))
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_after=Pt(0)
        i+=1
    doc.core_properties.title=TITLE
    doc.core_properties.subject="PlantShiftEval corrected methodological study"
    doc.core_properties.keywords="plant disease recognition, domain generalization, texture representation, DINOv2"
    OUTPUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size} bytes); figure files embedded: {sum(p.exists() for p in FIGURES.values())}/5")

if __name__=="__main__":main()
