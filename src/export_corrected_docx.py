"""Create an editable manuscript from the completed, audited Markdown study."""
from pathlib import Path
import argparse,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def inline(p,text):
    for token in re.split(r'(\*\*.*?\*\*|`.*?`|\*[^*]+\*)',text):
        run=p.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith(('`','*')) else token)
        run.bold=token.startswith('**');run.italic=token.startswith('*') and not token.startswith('**')
        run.font.name='Arial';run.font.color.rgb=RGBColor(0,0,0)

def table(doc,lines):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    n=len(rows[0]);widths=[6.5/n]*n
    if n==6:widths=[1.3,.65,.65,.55,1.8,1.55]
    elif n==4:widths=[1.8,1.55,1.55,1.6]
    elif n==3:widths=[1.8,2.0,2.7]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for row_idx,values in enumerate(rows):
        cells=t.rows[0].cells if row_idx==0 else t.add_row().cells
        for col,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            tcPr.append(borders);margin=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margin.append(e)
            tcPr.append(margin)
            fill='DDEBF3' if row_idx==0 else ('F6F8FA' if row_idx%2==0 else 'FFFFFF')
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),fill);tcPr.append(shading)
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if col==0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
            inline(p,value)
            for run in p.runs:run.font.size=Pt(9.5);run.bold=row_idx==0
        if row_idx==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        # Keep small tables intact and prevent individual rows splitting.
        row=t.rows[row_idx];row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if len(rows)<=6 and row_idx<len(rows)-1:
            for cell in cells:
                for para in cell.paragraphs:para.paragraph_format.keep_with_next=True
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    p=argparse.ArgumentParser();p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    md=(a.root/'manuscript/PlantShiftEval_corrected_manuscript.md').read_text()
    maintext,tail=md.split('## Figure legends',1);legendtext,refs=tail.split('## References',1)
    captions={int(m.group(1)):m.group(0) for m in re.finditer(r'\*\*Figure (\d+)\..*?(?=\n\n\*\*Figure|\Z)',legendtext,flags=re.S)}
    maintext+='\n## References'+refs
    lines=maintext.splitlines();doc=Document();sec=doc.sections[0]
    # The runtime default template may contain decorative paragraph borders.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(1)
    for name,size in [('Normal',11),('Title',18),('Subtitle',12),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
        fonts.set(qn('w:ascii'),'Arial');fonts.set(qn('w:hAnsi'),'Arial')
        fmt=style.paragraph_format;fmt.space_after=Pt(6);fmt.line_spacing=1.15
        if name.startswith('Heading'):fmt.space_before=Pt(12);fmt.keep_with_next=True
    doc.core_properties.title='Testing Texture Representation Hypotheses in Plant Disease Recognition';doc.core_properties.author='omixcore';doc.core_properties.subject='Corrected methodological study'
    title=doc.add_paragraph('Testing Texture Representation Hypotheses in Plant Disease Recognition','Title')
    doc.add_paragraph('Training Strategy Effects and Cross Domain Generalization Limits','Subtitle')
    i=1
    while i<len(lines):
        line=lines[i].strip()
        if not line or line.startswith('**Evidence status:'):i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(doc,block);continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));text=line.lstrip('#').strip()
            if text=='Results' and i+1<len(lines) and lines[i+1].strip()=='':pass
            doc.add_heading(re.sub(r'[-:;]',' ',text),level=min(3,max(1,level-1)))
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        else:
            para=doc.add_paragraph();para.paragraph_format.widow_control=True;inline(para,line)
        i+=1
    names=['Figure1_source_comparisons','Figure2_training_factorial','Figure3_component_ablations','Figure4_domain_limits','Figure5_target_access']
    for n,name in enumerate(names,1):
        if n in (1,3,5):doc.add_page_break()
        img=doc.add_paragraph();img.paragraph_format.keep_with_next=True;img.paragraph_format.space_after=Pt(3);img.alignment=WD_ALIGN_PARAGRAPH.CENTER
        img.add_run().add_picture(str(a.root/'report/corrected_summary'/f'{name}.png'),width=Inches(6.3))
        para=doc.add_paragraph();inline(para,captions[n].replace('\n',' '))
        for run in para.runs:run.font.size=Pt(9.5)
        para.paragraph_format.space_after=Pt(10)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    a.out.parent.mkdir(parents=True,exist_ok=True);doc.save(a.out);print(a.out)

if __name__=='__main__':main()
